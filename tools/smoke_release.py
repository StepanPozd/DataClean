"""Exercise the real Windows EXE through UI Automation, without embedded test hooks.

Run from a Windows desktop session with requirements-build.txt installed. All generated
inputs, outputs and reports stay in test-output/release/; none ship in the ZIP.
"""

import argparse
import csv
import hashlib
import json
import os
import subprocess
import sys
import time
from pathlib import Path
from uuid import uuid4

ROOT = Path(__file__).resolve().parents[1]


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--exe", type=Path, default=ROOT / "dist/DataClean/DataClean.exe")
    args = parser.parse_args()
    exe = args.exe.resolve()
    if not exe.is_relative_to(ROOT):
        raise ValueError("Test executable must be inside the project")
    work = ROOT / "test-output" / "release" / ("e2e-" + uuid4().hex[:8])
    work.mkdir(parents=True)
    data = work / "Кириллица и пробелы"
    data.mkdir()
    os.environ["TEMP"] = str(ROOT / ".cache/tmp")
    os.environ["TMP"] = os.environ["TEMP"]

    import ctypes
    from ctypes import wintypes

    import psutil
    from openpyxl import Workbook, load_workbook
    from pywinauto import Desktop

    user32 = ctypes.WinDLL("user32", use_last_error=True)
    user32.SendMessageW.argtypes = [wintypes.HWND, wintypes.UINT, wintypes.WPARAM, wintypes.LPARAM]
    user32.SendMessageW.restype = wintypes.LPARAM
    user32.GetParent.argtypes = [wintypes.HWND]
    user32.GetParent.restype = wintypes.HWND
    user32.GetDlgCtrlID.argtypes = [wintypes.HWND]
    user32.ScreenToClient.argtypes = [wintypes.HWND, ctypes.POINTER(wintypes.POINT)]
    user32.PostMessageW.argtypes = [wintypes.HWND, wintypes.UINT, wintypes.WPARAM, wintypes.LPARAM]

    headers = ["ID", "Name", "Token", "Amount"]
    first = ["00123", "  Анна  ", "NA", "0.12345678901234567890123456789"]
    long_id = "1234567890123456789012345678901234567890"
    rows = [
        first,
        ["00045", "Борис", "NULL", "2.500"],
        first,
        [""] * 4,
        [" "] * 4,
        [long_id, " Мария ", "N/A", "9007199254740993"],
    ]

    def write_csv(name, records):
        path = data / name
        with path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(headers)
            writer.writerows(records)
        return path

    source = write_csv("исходные.csv", rows)
    merge_sources = [
        write_csv(f"часть {index}.csv", [[value, "Имя", "NA", "1.00"]])
        for index, value in enumerate(["00010", "00011"])
    ]
    formula = write_csv("формула.csv", [["=1+1", "Текст", "NA", "1.00"]])
    long_text = write_csv("длинный текст.csv", [["00001", "Ж" * 40000, "NULL", "1.00"]])
    workbook = data / "книга с листами.xlsx"
    book = Workbook()
    book.active.title = "Первый"
    book.active.append(["City", "Value"])
    book.active.append(["Алматы", 1])
    sheet = book.create_sheet("Второй")
    sheet.append(["City", "Value"])
    sheet.append(["Астана", 2])
    sheet.append(["Казань", 3])
    book.save(workbook)
    inputs = [source, *merge_sources, formula, long_text, workbook]
    original_hashes = {str(path): hashlib.sha256(path.read_bytes()).hexdigest() for path in inputs}
    env = dict(os.environ)
    env.update(
        QT_ACCESSIBILITY="1",
        PYTHONHOME=str(work / "python-not-installed"),
        PYTHONPATH="",
        PATH=os.path.join(os.environ["SystemRoot"], "System32"),
    )
    env.pop("VIRTUAL_ENV", None)
    env.pop("QT_QPA_PLATFORM", None)
    process = subprocess.Popen([str(exe), str(source)], cwd=work, env=env)
    desktop = Desktop(backend="uia")
    native = Desktop(backend="win32")
    checks = {}
    main_handle = None
    report = {
        "exe": str(exe),
        "exe_sha256": hashlib.sha256(exe.read_bytes()).hexdigest(),
        "output": str(work),
        "checks": checks,
    }

    def until(function, description, timeout=35):
        started = time.monotonic()
        while time.monotonic() - started < timeout:
            if process.poll() is not None:
                raise RuntimeError(f"EXE exited during {description}: {process.returncode}")
            try:
                result = function()
                if result:
                    return result
            except Exception:
                pass
            time.sleep(0.15)
        raise TimeoutError(description)

    def window(fragment):
        def find():
            for item in native.windows(process=process.pid):
                if fragment in item.window_text():
                    return desktop.window(handle=item.handle)
            return None

        return until(find, "window: " + fragment)

    def button(parent, title=None, auto_id=None):
        if auto_id:
            return parent.child_window(auto_id=auto_id, control_type="Button")
        return parent.child_window(title=title, control_type="Button")

    def click(control):
        wrapper = control.wrapper_object() if hasattr(control, "wrapper_object") else control
        until(lambda: wrapper.is_enabled(), "enabled UI action")
        # Qt's checkable QPushButtons expose Toggle, which does not emit clicked.
        # Send mouse messages to this window only, without touching global focus.
        ancestor = wrapper
        while not ancestor.handle:
            ancestor = ancestor.parent()
        target = ancestor.handle
        print("UI action:", wrapper.window_text(), "target:", target, flush=True)
        if wrapper.handle and wrapper.element_info.class_name == "Button":
            assert user32.PostMessageW(target, 0x00F5, 0, 0)  # BM_CLICK
            time.sleep(0.2)
            return
        rect = wrapper.rectangle()
        point = wintypes.POINT((rect.left + rect.right) // 2, (rect.top + rect.bottom) // 2)
        assert user32.ScreenToClient(target, ctypes.byref(point))
        coordinates = (point.y << 16) | (point.x & 0xFFFF)
        assert user32.PostMessageW(target, 0x0201, 1, coordinates)
        assert user32.PostMessageW(target, 0x0202, 0, coordinates)
        time.sleep(0.2)

    def main_window():
        return desktop.window(handle=main_handle)

    def ready():
        until(lambda: button(main_window(), "Import File").is_enabled(), "worker completion")

    def row_count(count):
        ready()
        until(
            lambda: any(
                text.window_text().startswith(f"{count} rows") and "columns" in text.window_text()
                for text in main_window().descendants(control_type="Text")
            ),
            f"{count} rows",
        )

    def navigate(name):
        click(main_window().child_window(title=name, control_type="CheckBox"))
        until(
            lambda: (
                main_window()
                .child_window(auto_id="QApplication.MainWindow.root.topbar.pageTitle", control_type="Text")
                .window_text()
                == name
            ),
            "navigation: " + name,
        )

    def accept_csv():
        dialog = window("CSV import settings")
        click(button(dialog, "OK"))

    def accept_message(fragment, yes=False):
        dialog = window(fragment)

        def find_button():
            candidates = dialog.descendants(control_type="Button")
            wanted = {"Yes", "Да"} if yes else {"OK", "ОК"}
            return next((item for item in candidates if item.window_text().replace("&", "") in wanted), None)

        control = until(find_button, "message button: " + fragment)
        click(control)

    def file_dialog(fragment, paths, csv_output=False):
        # The native QFileDialog proxy appears before its Windows controls are ready.
        def usable():
            for item in native.windows(process=process.pid):
                if fragment in item.window_text():
                    spec = desktop.window(handle=item.handle)
                    for candidate in spec.descendants(control_type="Edit"):
                        identifier = candidate.element_info.automation_id
                        name = candidate.window_text().lower()
                        if (
                            identifier in {"1001", "1148", "1152"}
                            or "имя файла" in name
                            or "file name" in name
                        ):
                            return spec, candidate
            return None

        dialog, edit = until(usable, "native file dialog: " + fragment)
        if csv_output:
            combo = dialog.child_window(
                auto_id="FileTypeControlHost", control_type="ComboBox"
            ).wrapper_object()
            try:
                combo.select(1)
            except ValueError:
                # Windows' file-type combo applies the selection but does not expose
                # a SelectionItem pattern for pywinauto's post-selection index check.
                assert combo.selected_text() == "CSV formula-checked (*.csv)"
            time.sleep(0.5)
        value = str(paths[0]) if len(paths) == 1 else " ".join(f'"{path}"' for path in paths)
        assert edit.handle, "File-name edit must expose its native HWND"
        edit.iface_value.SetValue(value)
        identifier = user32.GetDlgCtrlID(edit.handle)
        user32.SendMessageW(
            user32.GetParent(edit.handle), 0x0111, (0x0300 << 16) | (identifier & 0xFFFF), edit.handle
        )  # EN_CHANGE
        time.sleep(0.2)
        assert edit.get_value() == value, edit.get_value()
        click(button(dialog, auto_id="1"))

    def import_file(path, csv_input=True):
        navigate("Dashboard")
        ready()
        click(button(main_window(), "Import File"))
        file_dialog("Import a dataset", [path])
        if csv_input:
            accept_csv()
        else:
            dialog = window("Select worksheet")
            combo = dialog.child_window(control_type="ComboBox").wrapper_object()
            click(combo)
            item = until(
                lambda: next(
                    (
                        item
                        for item in combo.descendants(control_type="ListItem")
                        if item.window_text() == "Второй"
                    ),
                    None,
                ),
                "second worksheet item",
            )
            click(item)
            until(lambda: combo.selected_text() == "Второй", "worksheet selection")
            click(button(dialog, "OK"))
        until(lambda: path.name in native.window(handle=main_handle).window_text(), "import title")
        ready()

    def export(path, error=None):
        click(button(main_window(), "Export"))
        file_dialog("Export dataset", [path], csv_output=path.suffix == ".csv")
        if path.suffix == ".csv":
            accept_message("CSV export settings", yes=True)
        if error:
            dialog = window("DataClean · Please check")
            texts = [item.window_text() for item in dialog.descendants(control_type="Text")]
            assert any(error in text for text in texts), texts
            accept_message("DataClean · Please check")
            assert not path.exists()
        else:
            accept_message("Export complete")
            assert path.is_file()
        ready()

    def csv_records(path):
        with path.open(encoding="utf-8-sig", newline="") as handle:
            return list(csv.reader(handle))

    try:
        accept_csv()
        main_handle = window("DataClean — исходные.csv").handle
        row_count(6)
        checks["main_window_csv_import_cyrillic_paths"] = True
        modules = [item.path for item in psutil.Process(process.pid).memory_maps()]
        bundled = [path for path in modules if Path(path).name.lower().startswith(("python3", "qt6"))]
        assert bundled and all(Path(path).is_relative_to(exe.parent) for path in bundled), bundled
        assert not any(".venv" in path.lower() for path in modules)
        report["loaded_runtime_modules"] = bundled
        report["loaded_icu"] = [path for path in modules if "icu" in Path(path).name.lower()]
        checks["bundled_python_and_qt_no_venv_modules"] = True
        navigate("Data Editor")
        click(button(main_window(), "Remove Duplicates"))
        row_count(5)
        checks["remove_duplicates"] = True
        click(button(main_window(), "Remove Empty Rows"))
        row_count(3)
        checks["remove_empty_rows"] = True
        click(button(main_window(), "Trim Whitespace"))
        ready()
        cleaned_csv = data / "очищенные.csv"
        export(cleaned_csv)
        clean = csv_records(cleaned_csv)
        assert [row[0] for row in clean[1:]] == ["00123", "00045", long_id]
        assert [row[1] for row in clean[1:]] == ["Анна", "Борис", "Мария"]
        assert clean[1][2:] == ["NA", first[3]]
        checks["trim_csv_export_identifiers_and_precision"] = True
        click(button(main_window(), "Undo Last Action"))
        ready()
        undo_csv = data / "после отмены.csv"
        export(undo_csv)
        assert csv_records(undo_csv)[1][1] == "  Анна  "
        checks["undo"] = True
        click(button(main_window(), "Reset to Original"))
        accept_message("Reset to original?", yes=True)
        row_count(6)
        checks["reset"] = True
        click(button(main_window(), "Undo Last Action"))
        row_count(3)
        click(button(main_window(), "Trim Whitespace"))
        ready()
        cleaned_xlsx = data / "очищенные.xlsx"
        export(cleaned_xlsx)
        book = load_workbook(cleaned_xlsx, read_only=True)
        assert book.active["A2"].value == "00123" and book.active["A2"].data_type == "s"
        assert book.active["A4"].value == long_id
        book.close()
        checks["xlsx_export_literal_identifiers"] = True
        import_file(workbook, csv_input=False)
        row_count(2)
        result = data / "выбранный лист.csv"
        export(result)
        assert csv_records(result)[1][0] == "Астана"
        checks["xlsx_import_selected_sheet"] = True
        navigate("Merge Files")
        click(button(main_window(), "Add files"))
        file_dialog("Select files to combine", merge_sources)
        for _ in merge_sources:
            accept_csv()
            time.sleep(0.3)
        ready()
        click(button(main_window(), "Merge files"))
        row_count(2)
        merged = data / "объединённые.csv"
        export(merged)
        assert [row[0] for row in csv_records(merged)[1:]] == ["00010", "00011"]
        checks["merge_csv_files"] = True
        import_file(formula)
        export(data / "заблокированная формула.csv", error="formula")
        checks["csv_formula_protection"] = True
        literal = data / "формула как текст.xlsx"
        export(literal)
        book = load_workbook(literal, read_only=True, data_only=False)
        assert book.active["A2"].value == "=1+1" and book.active["A2"].data_type == "s"
        book.close()
        checks["xlsx_formula_text_literal"] = True
        import_file(long_text)
        export(data / "слишком длинный.xlsx", error="32,767")
        checks["xlsx_long_text_blocked"] = True
        complete = data / "полный длинный текст.csv"
        export(complete)
        assert len(csv_records(complete)[1][1]) == 40000
        checks["csv_long_text_complete"] = True
        navigate("History")
        assert all(
            hashlib.sha256(Path(path).read_bytes()).hexdigest() == expected
            for path, expected in original_hashes.items()
        )
        checks["all_six_input_files_unchanged"] = True
        main_window().close()
        assert process.wait(timeout=20) == 0
        checks["clean_shutdown"] = True
        report["success"] = True
    except Exception as error:
        report["success"] = False
        report["error"] = f"{type(error).__name__}: {error}"
        raise
    finally:
        if process.poll() is None:
            process.terminate()
            process.wait(timeout=10)
        (work / "results.json").write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")
        print(json.dumps(report, indent=2, ensure_ascii=True), flush=True)


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    main()
