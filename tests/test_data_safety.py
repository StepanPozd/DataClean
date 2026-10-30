"""Data-loss and spreadsheet-formula regressions, separate from the original suite."""

import csv
import hashlib
import warnings

import pandas as pd
import pytest
from openpyxl import load_workbook
from pandas.testing import assert_frame_equal

from app.core import file_handler
from app.core.file_handler import ExportValidationError, export_file, read_file


@pytest.mark.parametrize("length", [32766, 32767])
def test_xlsx_text_at_limit_is_written_without_warning_or_truncation(tmp_path, length):
    text = "Я" * length
    path = tmp_path / "Граница.xlsx"
    with warnings.catch_warnings():
        warnings.simplefilter("error")
        export_file(pd.DataFrame({"Text": [text]}), path)
    book = load_workbook(path, read_only=True)
    try:
        assert book.active["A2"].value == text
    finally:
        book.close()


@pytest.mark.parametrize("length", [32768, 40000])
@pytest.mark.parametrize("existing", [False, True])
def test_xlsx_over_limit_is_rejected_before_any_file_is_created(tmp_path, monkeypatch, length, existing):
    path = tmp_path / "existing.xlsx"
    if existing:
        path.write_bytes(b"previous output")

    def no_temporary_files(*args, **kwargs):
        pytest.fail("Validation must run before opening a temporary output")

    monkeypatch.setattr(file_handler.tempfile, "mkstemp", no_temporary_files)
    frame = pd.DataFrame({"Name": ["safe"], "Notes": ["=Т" + "я" * (length - 2)]})
    original = frame.copy(deep=True)
    with pytest.raises(ExportValidationError, match="row 2, column 2") as error:
        export_file(frame, path, overwrite=existing)
    assert "32,767" in str(error.value)
    assert "CSV" in str(error.value)
    assert "not" in str(error.value)
    assert_frame_equal(frame, original)
    assert path.read_bytes() == b"previous output" if existing else not path.exists()


def test_xlsx_long_header_is_also_validated_in_a_header_only_dataset(tmp_path):
    path = tmp_path / "header.xlsx"
    with pytest.raises(ExportValidationError, match="header, column 1"):
        export_file(pd.DataFrame(columns=["H" * 32768]), path)
    assert not path.exists()
    assert not list(tmp_path.glob(".dataclean-*"))


def test_csv_alternative_keeps_the_entire_long_value(tmp_path):
    value = "Данные " * 7000
    frame = pd.DataFrame({"Notes": [value]})
    path = tmp_path / "long.csv"
    export_file(frame, path)
    assert read_file(path, csv_mode="text").frame.iat[0, 0] == value


@pytest.mark.parametrize("encoding", ["utf-8-sig", "cp1251", "utf-16"])
def test_text_import_preserves_identifiers_null_tokens_and_precision(tmp_path, encoding):
    values = [
        "00123",
        "00045",
        "NA",
        "NULL",
        "NaN",
        "N/A",
        "null",
        "",
        "9007199254740993",
        "1234567890123456789012345678901234567890",
        "0.123456789012345678901234567890",
        "1.2300",
        "1e30",
        " 00123 ",
        "Кириллица",
    ]
    path = tmp_path / "Идентификаторы.csv"
    with path.open("w", encoding=encoding, newline="") as handle:
        writer = csv.writer(handle, delimiter=";")
        writer.writerow(["Value", "Other"])
        writer.writerows((value, "anchor") for value in values)
    before = hashlib.sha256(path.read_bytes()).hexdigest()
    loaded = read_file(path, encoding=encoding, csv_mode="text")
    assert loaded.frame["Value"].tolist() == values
    assert all(isinstance(value, str) for value in loaded.frame["Value"])
    assert not loaded.frame.isna().any().any()
    assert loaded.csv_mode == "text"
    assert hashlib.sha256(path.read_bytes()).hexdigest() == before


def test_text_mode_preserves_empty_records_and_short_rows(tmp_path):
    path = tmp_path / "empty.csv"
    path.write_text('A,B\n00123,NA\n\n00045\n"",NULL\n', encoding="utf-8")
    frame = read_file(path, csv_mode="text").frame
    assert frame.values.tolist() == [["00123", "NA"], ["", ""], ["00045", ""], ["", "NULL"]]


def test_auto_interpretation_remains_explicit_and_compatible(tmp_path):
    path = tmp_path / "auto.csv"
    path.write_text("ID,Token,Amount\n00123,NA,1.25\n00045,NULL,2.50\n", encoding="utf-8")
    text = read_file(path, csv_mode="text").frame
    automatic = read_file(path, csv_mode="auto").frame
    assert text["ID"].tolist() == ["00123", "00045"]
    assert text["Token"].tolist() == ["NA", "NULL"]
    assert text["Amount"].tolist() == ["1.25", "2.50"]
    assert automatic["ID"].tolist() == [123, 45]
    assert automatic["Token"].isna().all()
    assert automatic["Amount"].tolist() == [1.25, 2.5]


def test_invalid_import_interpretation_is_rejected(tmp_path):
    path = tmp_path / "data.csv"
    path.write_text("A\n00123\n", encoding="utf-8")
    with pytest.raises(ValueError, match="interpretation"):
        read_file(path, csv_mode="guess")


@pytest.mark.parametrize(
    "value",
    [
        "=1+1",
        "+SUM(A1:A2)",
        "-1+2",
        "@SUM(A1:A2)",
        " =1+1",
        "\ttext",
        "\rtext",
        "\ntext",
        " \t=1+1",
        "\ufeff=1+1",
        "\u200b=1+1",
        "'=1+1",
        "＝1+1",
        "＋1+1",
        "－1+1",
        "＠SUM(A1:A2)",
        "-12",
    ],
)
def test_default_csv_mode_blocks_formula_candidates_without_creating_output(tmp_path, value):
    frame = pd.DataFrame({"Text": [value]})
    original = frame.copy(deep=True)
    path = tmp_path / "blocked.csv"
    with pytest.raises(ExportValidationError, match="formula") as error:
        export_file(frame, path)
    assert "XLSX" in str(error.value)
    assert "row 2, column 1" in str(error.value)
    assert not path.exists()
    assert not list(tmp_path.glob(".dataclean-*"))
    assert_frame_equal(frame, original)


def test_formula_like_csv_headers_are_checked_and_previous_output_is_preserved(tmp_path):
    path = tmp_path / "header.csv"
    path.write_bytes(b"previous data")
    with pytest.raises(ExportValidationError, match="header, column 1"):
        export_file(pd.DataFrame(columns=["=1+1"]), path, overwrite=True)
    assert path.read_bytes() == b"previous data"


def test_formula_checked_csv_keeps_safe_values_and_quotes_embedded_separators(tmp_path):
    values = ["00123", "NA", "NULL", "a=b", "name@example.test", 'safe,"quoted"\n=still inside one cell']
    frame = pd.DataFrame({"Text": values, "Number": [-1, -2, 3, 0, 1.25, -0.5]})
    path = tmp_path / "safe.csv"
    export_file(frame, path)
    with path.open(encoding="utf-8-sig", newline="") as handle:
        records = list(csv.reader(handle))
    assert [record[0] for record in records[1:]] == values
    assert records[1][1] == "-1.0"
    assert path.read_bytes().startswith(b"\xef\xbb\xbf")


def test_raw_csv_preserves_formula_text_headers_quotes_and_newlines_exactly(tmp_path):
    values = ["=1+1", "+SUM(A1:A2)", "-2", "@name", "\t=1", "  =1", 'a,"b"\r\nc', "00123"]
    frame = pd.DataFrame({"=Header": values})
    original = frame.copy(deep=True)
    path = tmp_path / "raw.csv"
    export_file(frame, path, csv_mode="raw")
    with path.open(encoding="utf-8-sig", newline="") as handle:
        records = list(csv.reader(handle))
    assert records == [["=Header"], *[[value] for value in values]]
    assert_frame_equal(frame, original)


def test_invalid_csv_export_mode_fails_before_creating_output(tmp_path):
    path = tmp_path / "invalid.csv"
    with pytest.raises(ExportValidationError, match="mode"):
        export_file(pd.DataFrame({"A": [1]}), path, csv_mode="escape")
    assert not path.exists()


def test_safe_csv_import_to_xlsx_keeps_text_types_and_formula_text(tmp_path):
    values = [
        "00123",
        "00045",
        "NA",
        "NULL",
        "123456789012345678901234567890",
        "0.12345678901234567890123456789",
        "=1+1",
        "+1+1",
        "-2",
        "@name",
    ]
    source = tmp_path / "source.csv"
    with source.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["Text"])
        writer.writerows([value] for value in values)
    before = source.read_bytes()
    target = tmp_path / "literal.xlsx"
    export_file(read_file(source, csv_mode="text").frame, target)
    book = load_workbook(target, read_only=True, data_only=False)
    try:
        cells = [row[0] for row in book.active.iter_rows(min_row=2)]
        assert [cell.value for cell in cells] == values
        assert all(cell.data_type == "s" for cell in cells)
    finally:
        book.close()
    assert source.read_bytes() == before
