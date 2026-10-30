"""Local file I/O with explicit formats, source protection and atomic export."""

import codecs
import csv
import io
import os
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Literal

import pandas as pd
from charset_normalizer import from_bytes
from openpyxl import load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

CsvImportMode = Literal["auto", "text"]
CsvExportMode = Literal["spreadsheet", "raw"]
EXCEL_CELL_CHARACTERS = 32_767


class ExportValidationError(ValueError):
    """An export would lose data or violate the chosen safety policy."""


class FileReadError(ValueError):
    """A readable, actionable import failure."""


class EncodingReadError(FileReadError):
    """CSV could not be decoded; the interface can offer manual encoding."""


@dataclass(frozen=True)
class LoadedFile:
    path: Path
    frame: pd.DataFrame
    encoding: str | None = None
    sheet: str | None = None
    csv_mode: CsvImportMode | None = None


def canonical(path: str | Path) -> str:
    return os.path.normcase(str(Path(path).resolve()))


def excel_sheets(path: str | Path) -> list[str]:
    try:
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            return book.sheetnames
        finally:
            book.close()
    except Exception as error:
        raise FileReadError(f"Cannot open Excel workbook: {error}") from error


def detect_encoding(raw: bytes) -> str:
    if raw.startswith(codecs.BOM_UTF8):
        return "utf-8-sig"
    if raw.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return "utf-16"
    try:
        raw.decode("utf-8")
        return "utf-8"
    except UnicodeDecodeError:
        candidate = from_bytes(raw, cp_isolation=["cp1251", "cp1252", "cp866", "koi8_r", "iso8859_1"]).best()
        if candidate is None:
            raise EncodingReadError("CSV encoding was not recognized. Choose an encoding manually.")
        return candidate.encoding


def validate_headers(headers: list) -> None:
    if not headers or all(value is None or str(value).strip() == "" for value in headers):
        raise FileReadError("The file is empty or has no header row.")
    if any(value is None or not str(value).strip() for value in headers):
        raise FileReadError("Every column must have a non-empty name in the first row.")
    names = [str(value) for value in headers]
    if len(set(names)) != len(names):
        raise FileReadError("Duplicate column names found. Give each column a unique name.")


def read_file(
    path: str | Path,
    *,
    encoding: str | None = None,
    sheet: str | None = None,
    csv_mode: CsvImportMode = "auto",
) -> LoadedFile:
    source = Path(path).resolve()
    suffix = source.suffix.lower()
    if suffix not in {".csv", ".xlsx"}:
        raise FileReadError("Choose a CSV or XLSX file.")
    try:
        if suffix == ".csv":
            if csv_mode not in {"auto", "text"}:
                raise FileReadError("Choose CSV interpretation 'text' or 'auto'.")
            raw = source.read_bytes()
            selected = encoding or detect_encoding(raw)
            try:
                content = raw.decode(selected)
            except (UnicodeError, LookupError) as error:
                raise EncodingReadError(
                    f"Cannot decode CSV as {selected}. Choose another encoding."
                ) from error
            if "\x00" in content:
                raise EncodingReadError("CSV contains unreadable characters. Try UTF-16.")
            try:
                dialect = csv.Sniffer().sniff(content[:65536], delimiters=",;\t|")
                delimiter = dialect.delimiter
            except csv.Error:
                first_line = content.splitlines()[0] if content.splitlines() else ""
                delimiter = (
                    max(",;\t|", key=first_line.count) if any(char in first_line for char in ",;\t|") else ","
                )
            records = csv.reader(io.StringIO(content), delimiter=delimiter, strict=True)
            header = next(records, [])
            validate_headers(header)
            # pandas can otherwise interpret an extra first field as an implicit index,
            # silently losing that field on an index-free export.
            for record in records:
                if len(record) > len(header):
                    raise FileReadError(
                        f"CSV line {records.line_num} has {len(record)} fields; the header has "
                        f"{len(header)}. Check separators and quote text containing delimiters."
                    )
            # Text mode preserves tokens and precision, including empty fields.
            # The legacy auto mode remains explicit for callers and UI users who need numeric data.
            options = (
                {"dtype": str, "na_filter": False, "keep_default_na": False} if csv_mode == "text" else {}
            )
            frame = pd.read_csv(io.StringIO(content), sep=delimiter, skip_blank_lines=False, **options)
            return LoadedFile(source, frame, selected, csv_mode=csv_mode)
        book = load_workbook(source, read_only=True, data_only=True)
        try:
            selected_sheet = sheet if sheet is not None else book.sheetnames[0]
            if selected_sheet not in book.sheetnames:
                raise FileReadError(f"Worksheet '{selected_sheet}' does not exist.")
            header = list(next(book[selected_sheet].iter_rows(values_only=True), ()))
            validate_headers(header)
        finally:
            book.close()
        frame = pd.read_excel(source, sheet_name=selected_sheet, engine="openpyxl")
        frame.columns = frame.columns.map(str)
        return LoadedFile(source, frame, sheet=selected_sheet)
    except FileReadError:
        raise
    except Exception as error:
        raise FileReadError(f"Cannot read {source.name}: {error}") from error


def _export_values(frame: pd.DataFrame):
    """Match the flat, index-free table written by our exporters, including headers."""
    if isinstance(frame.columns, pd.MultiIndex):
        raise ExportValidationError("Export requires a table with a single header row.")
    for column, value in enumerate(frame.columns, 1):
        yield f"header, column {column}", str(value)
    for row, values in enumerate(frame.itertuples(index=False, name=None), 2):
        for column, value in enumerate(values, 1):
            yield f"row {row}, column {column}", value


def _validate_excel(frame: pd.DataFrame) -> None:
    if len(frame) > 1_048_575 or len(frame.columns) > 16_384:
        raise ExportValidationError("The dataset exceeds Excel's row or column limit. Export to CSV instead.")
    for location, value in _export_values(frame):
        # pandas also serializes object values (e.g. Decimal) to text. Check their
        # string representation before any writer has a chance to truncate it.
        text = value if isinstance(value, str) else str(value)
        if len(text) > EXCEL_CELL_CHARACTERS:
            raise ExportValidationError(
                f"Cannot export XLSX: {location} contains {len(text):,} characters. "
                f"Excel allows at most {EXCEL_CELL_CHARACTERS:,} characters per cell. "
                "No data was written or truncated. Export to CSV to keep the full value."
            )
        if ILLEGAL_CHARACTERS_RE.search(text):
            raise ExportValidationError(
                f"Cannot export XLSX: {location} contains control characters Excel cannot store. "
                "Export to CSV instead."
            )


def _formula_candidate(value: str) -> bool:
    # Conservative: quoted numeric strings such as '-12' are intentionally flagged.
    # Examine leading whitespace/control characters, apostrophes and full-width signs
    # without ever stripping, escaping or otherwise changing the exported value.
    for character in value:
        if character in "\t\r\n":
            return True
        if character.isspace() or character in "\ufeff\u200b'":
            continue
        return character in "=+-@＝＋－＠"
    return False


def _validate_csv(frame: pd.DataFrame, mode: CsvExportMode) -> None:
    if mode not in {"spreadsheet", "raw"}:
        raise ExportValidationError("Choose CSV export mode 'spreadsheet' or 'raw'.")
    if mode == "spreadsheet":
        for location, value in _export_values(frame):
            if isinstance(value, str) and _formula_candidate(value):
                raise ExportValidationError(
                    f"CSV export blocked: {location} could be interpreted as a formula by a spreadsheet. "
                    "Values were not escaped or changed. Use XLSX to store formula-like text literally, "
                    "or explicitly choose 'CSV exact values (unprotected)' for data exchange. "
                    "Do not open unprotected CSV directly in a spreadsheet."
                )


def export_file(
    frame: pd.DataFrame,
    path: str | Path,
    *,
    protected_paths: set[str] | None = None,
    overwrite: bool = False,
    csv_mode: CsvExportMode = "spreadsheet",
) -> Path:
    target = Path(path).resolve()
    if target.suffix.lower() not in {".csv", ".xlsx"}:
        raise ValueError("Use a .csv or .xlsx extension.")
    for protected in protected_paths or set():
        if canonical(target) == canonical(protected) or (
            target.exists() and Path(protected).exists() and os.path.samefile(target, protected)
        ):
            raise ValueError("Source files are protected. Export using a new filename.")
    if target.exists() and not overwrite:
        raise FileExistsError(
            "A file with this name already exists. Choose another name or confirm overwrite."
        )
    # Complete validation precedes even creation of the temporary output file.
    if target.suffix.lower() == ".xlsx":
        _validate_excel(frame)
    else:
        _validate_csv(frame, csv_mode)
    handle, temporary = tempfile.mkstemp(prefix=".dataclean-", suffix=target.suffix, dir=target.parent)
    os.close(handle)
    try:
        if target.suffix.lower() == ".csv":
            frame.to_csv(temporary, index=False, encoding="utf-8-sig")
        else:
            # Write strings literally: a value beginning with '=' is data, not a formula.
            with pd.ExcelWriter(temporary, engine="openpyxl") as writer:
                frame.to_excel(writer, index=False, sheet_name="Data")
                worksheet = writer.sheets["Data"]
                for row in worksheet:
                    for cell in row:
                        if cell.data_type == "f":
                            cell.data_type = "s"
                worksheet.freeze_panes = "A2"
                worksheet.auto_filter.ref = worksheet.dimensions
        if overwrite:
            os.replace(temporary, target)
        elif os.name == "nt":
            os.rename(temporary, target)  # Windows refuses an existing destination atomically.
        else:
            os.link(temporary, target)
            os.unlink(temporary)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)
    return target
